"""
AI Career Coach - Document Processing Utilities
Handles PDF/DOCX parsing and export.
"""
import os
import uuid
from pathlib import Path
from config import Config


def allowed_file(filename):
    """Check if file extension is allowed."""
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in Config.ALLOWED_EXTENSIONS


def save_upload(file_obj, purpose='general'):
    """Save an uploaded file and return metadata."""
    if not file_obj or not file_obj.filename:
        return None
    
    if not allowed_file(file_obj.filename):
        return None
    
    Config.ensure_dirs()
    ext = file_obj.filename.rsplit('.', 1)[1].lower()
    filename = f"{uuid.uuid4().hex}.{ext}"
    filepath = Config.UPLOAD_FOLDER / filename
    file_obj.save(str(filepath))
    
    file_size = filepath.stat().st_size
    extracted_text = extract_text(filepath, ext)
    
    return {
        'filename': filename,
        'original_filename': file_obj.filename,
        'file_type': ext,
        'file_size': file_size,
        'extracted_text': extracted_text,
        'purpose': purpose
    }


def extract_text(filepath, file_type=None):
    """Extract text content from a file."""
    filepath = Path(filepath)
    if file_type is None:
        file_type = filepath.suffix.lstrip('.').lower()
    
    try:
        if file_type == 'pdf':
            return extract_pdf_text(filepath)
        elif file_type in ('docx', 'doc'):
            return extract_docx_text(filepath)
        elif file_type == 'txt':
            return filepath.read_text(encoding='utf-8')
        else:
            return ''
    except Exception as e:
        print(f"Text extraction error: {e}")
        return f"[Error extracting text: {str(e)}]"


def extract_pdf_text(filepath):
    """Extract text from PDF."""
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(str(filepath))
        text_parts = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
        return '\n'.join(text_parts)
    except ImportError:
        return "[PDF parsing requires PyPDF2. Install with: pip install PyPDF2]"
    except Exception as e:
        return f"[Error reading PDF: {str(e)}]"


def extract_docx_text(filepath):
    """Extract text from DOCX."""
    try:
        from docx import Document
        doc = Document(str(filepath))
        text_parts = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)
        return '\n'.join(text_parts)
    except ImportError:
        return "[DOCX parsing requires python-docx. Install with: pip install python-docx]"
    except Exception as e:
        return f"[Error reading DOCX: {str(e)}]"


def _xml_esc(val):
    """Escape text for ReportLab XML Paragraph parser."""
    if val is None:
        return ''
    return (
        str(val)
        .replace('&', '&amp;')
        .replace('<', '&lt;')
        .replace('>', '&gt;')
    )


def _prepare_cv_data(cv_data):
    """Ensure contact_info, sections, and summary are normalized."""
    try:
        from database import _normalize_cv_data
        contact, sections, summary = _normalize_cv_data(cv_data)
    except Exception:
        contact = cv_data.get('contact_info', {})
        sections = cv_data.get('sections', [])
        summary = cv_data.get('professional_summary', '')
    return contact, sections, summary


def _register_unicode_font():
    """Register a Unicode-capable TTF font (Arial/Tahoma on Windows) for PDF generation."""
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        font_candidates = [
            ('CVFont', 'CVFont-Bold', r'C:\Windows\Fonts\arial.ttf', r'C:\Windows\Fonts\arialbd.ttf'),
            ('CVFont', 'CVFont-Bold', r'C:\Windows\Fonts\tahoma.ttf', r'C:\Windows\Fonts\tahomabd.ttf'),
        ]
        for reg_name, bold_name, reg_path, bold_path in font_candidates:
            if Path(reg_path).exists():
                pdfmetrics.registerFont(TTFont(reg_name, reg_path))
                if Path(bold_path).exists():
                    pdfmetrics.registerFont(TTFont(bold_name, bold_path))
                else:
                    pdfmetrics.registerFont(TTFont(bold_name, reg_path))
                return reg_name, bold_name
    except Exception:
        pass
    return 'Helvetica', 'Helvetica-Bold'


def generate_cv_pdf(cv_data):
    """Generate a PDF from CV data."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib.colors import HexColor
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
        from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
        import io
        
        buffer = io.BytesIO()
        is_rtl = cv_data.get('language', 'en') == 'ar'
        alignment = TA_RIGHT if is_rtl else TA_LEFT
        template = cv_data.get('template', 'professional')
        font_reg, font_bold = _register_unicode_font()
        
        accent_map = {
            'professional': ('#1a365d', '#2d3748'),
            'modern':       ('#4338ca', '#312e81'),
            'simple':       ('#111827', '#374151'),
        }
        primary_hex, heading_hex = accent_map.get(template, accent_map['professional'])

        doc = SimpleDocTemplate(buffer, pagesize=A4,
                               rightMargin=1.5*cm, leftMargin=1.5*cm,
                               topMargin=1.5*cm, bottomMargin=1.5*cm)
        
        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle(
            'CVTitle', parent=styles['Title'],
            fontName=font_bold,
            fontSize=20, spaceAfter=6, alignment=TA_CENTER,
            textColor=HexColor(primary_hex)
        )
        
        section_style = ParagraphStyle(
            'CVSection', parent=styles['Heading2'],
            fontName=font_bold,
            fontSize=13, spaceAfter=6, spaceBefore=12,
            textColor=HexColor(heading_hex),
            borderWidth=0, borderPadding=0,
            alignment=alignment
        )
        
        body_style = ParagraphStyle(
            'CVBody', parent=styles['Normal'],
            fontName=font_reg,
            fontSize=10, spaceAfter=4,
            alignment=alignment,
            leading=14
        )
        
        contact_style = ParagraphStyle(
            'CVContact', parent=styles['Normal'],
            fontName=font_reg,
            fontSize=9, alignment=TA_CENTER,
            textColor=HexColor('#4a5568')
        )
        
        elements = []
        contact, sections, summary = _prepare_cv_data(cv_data)
        
        name = contact.get('name') or cv_data.get('title', 'Untitled CV')
        elements.append(Paragraph(_xml_esc(name), title_style))
        
        contact_parts = []
        for k in ('email', 'phone', 'location', 'linkedin'):
            if contact.get(k):
                contact_parts.append(_xml_esc(contact[k]))
        
        if contact_parts:
            elements.append(Paragraph(' | '.join(contact_parts), contact_style))
        
        elements.append(Spacer(1, 12))
        elements.append(HRFlowable(width="100%", thickness=1, color=HexColor('#cbd5e0')))
        elements.append(Spacer(1, 8))
        
        # Professional summary
        if summary:
            elements.append(Paragraph('Professional Summary' if not is_rtl else 'الملخص المهني', section_style))
            elements.append(Paragraph(_xml_esc(summary), body_style))
            elements.append(Spacer(1, 6))
        
        # Sections
        for section in sections:
            section_title = section.get('title', '')
            if section_title:
                elements.append(HRFlowable(width="100%", thickness=0.5, color=HexColor('#e2e8f0')))
                elements.append(Paragraph(_xml_esc(section_title), section_style))
            
            items = section.get('items', [])
            for item in items:
                if isinstance(item, dict):
                    header = item.get('title', item.get('degree', item.get('name', '')))
                    subtitle = item.get('company', item.get('institution', item.get('issuer', '')))
                    date = item.get('date', item.get('dates', ''))
                    tech = item.get('technologies', '')
                    link = item.get('link', '')
                    
                    if header:
                        header_text = f"<b>{_xml_esc(header)}</b>"
                        if subtitle:
                            header_text += f" — {_xml_esc(subtitle)}"
                        if date:
                            header_text += f"  <i>({_xml_esc(date)})</i>"
                        if tech:
                            header_text += f"  [{_xml_esc(tech)}]"
                        if link:
                            header_text += f"  ({_xml_esc(link)})"
                        elements.append(Paragraph(header_text, body_style))
                    
                    description = item.get('description', '')
                    if description:
                        for line in str(description).split('\n'):
                            line = line.strip()
                            if line:
                                clean = _xml_esc(line.lstrip('•-* '))
                                elements.append(Paragraph(f"    • {clean}", body_style))
                    
                    bullets = item.get('bullets', [])
                    for bullet in bullets:
                        elements.append(Paragraph(f"    • {_xml_esc(bullet)}", body_style))
                    
                elif isinstance(item, str):
                    elements.append(Paragraph(f"• {_xml_esc(item)}", body_style))
            
            elements.append(Spacer(1, 4))
        
        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()
    
    except ImportError:
        return None
    except Exception as e:
        print(f"PDF generation error: {e}")
        return None


def generate_cv_docx(cv_data):
    """Generate a DOCX from CV data."""
    try:
        from docx import Document
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        import io
        
        doc = Document()
        is_rtl = cv_data.get('language', 'en') == 'ar'
        
        # Set margins
        for section in doc.sections:
            section.top_margin = Cm(1.5)
            section.bottom_margin = Cm(1.5)
            section.left_margin = Cm(2)
            section.right_margin = Cm(2)
        
        style = doc.styles['Normal']
        style.font.size = Pt(10)
        style.font.name = 'Calibri'
        
        contact, sections, summary = _prepare_cv_data(cv_data)
        
        name = contact.get('name') or cv_data.get('title', '')
        if name:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(name))
            run.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(26, 54, 93)
        
        contact_parts = []
        for field in ['email', 'phone', 'location', 'linkedin']:
            if contact.get(field):
                contact_parts.append(str(contact[field]))
        if contact_parts:
            p = doc.add_paragraph(' | '.join(contact_parts))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(74, 85, 104)
        
        # Summary
        if summary:
            h = doc.add_heading('Professional Summary' if not is_rtl else 'الملخص المهني', level=2)
            for run in h.runs:
                run.font.color.rgb = RGBColor(45, 55, 72)
            p_sum = doc.add_paragraph(str(summary))
            if is_rtl:
                p_sum.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        # Sections
        for section in sections:
            title = section.get('title', '')
            if title:
                h = doc.add_heading(str(title), level=2)
                for run in h.runs:
                    run.font.color.rgb = RGBColor(45, 55, 72)
            
            items = section.get('items', [])
            for item in items:
                if isinstance(item, dict):
                    header = item.get('title', item.get('degree', item.get('name', '')))
                    subtitle = item.get('company', item.get('institution', item.get('issuer', '')))
                    date = item.get('date', item.get('dates', ''))
                    tech = item.get('technologies', '')
                    link = item.get('link', '')
                    
                    if header:
                        p = doc.add_paragraph()
                        if is_rtl:
                            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        run = p.add_run(str(header))
                        run.bold = True
                        if subtitle:
                            p.add_run(f' — {subtitle}')
                        if date:
                            run2 = p.add_run(f'  ({date})')
                            run2.italic = True
                        if tech:
                            p.add_run(f'  [{tech}]')
                        if link:
                            p.add_run(f'  ({link})')
                    
                    desc = item.get('description', '')
                    if desc:
                        for line in str(desc).split('\n'):
                            line = line.strip()
                            if line:
                                clean = line.lstrip('•-* ')
                                doc.add_paragraph(clean, style='List Bullet')
                    
                    for bullet in item.get('bullets', []):
                        doc.add_paragraph(str(bullet), style='List Bullet')
                        
                elif isinstance(item, str):
                    doc.add_paragraph(str(item), style='List Bullet')
        
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
    
    except ImportError:
        return None
    except Exception as e:
        print(f"DOCX generation error: {e}")
        return None
